"""Document loaders.

Each loader takes bytes plus a filename and returns a `ParsedDocument`: the
structure the format actually declares, not a wall of text. Headings come from
heading styles, from `<h1>`, from `#`, or - in a PDF, which declares nothing -
from the type size the document itself uses for its headings. Nothing here
guesses structure out of flat prose: that is how numbered list items ended up
being treated as section headings, shredding a page into ten fragments.
"""

from __future__ import annotations

import io
import json
from collections.abc import Callable
from typing import Any, Final

import structlog

from zk2.sources.chunking import markdown_heading
from zk2.sources.parsed import ParsedDocument, Segment
from zk2.sources.text_clean import clean_text

logger = structlog.get_logger()

SUPPORTED_EXTENSIONS: Final[frozenset[str]] = frozenset(
    {"pdf", "txt", "md", "json", "docx", "xlsx", "html", "htm"}
)

#: Sections that hold a table of contents. Their text is page numbers and dot
#: leaders: it matches every query that quotes a section title and answers
#: none of them.
_CONTENTS_TITLES: Final[frozenset[str]] = frozenset(
    {"table of contents", "contents", "оглавление", "содержание"}
)

#: A heading is short. A paragraph that happens to be set large is not one.
_MAX_HEADING_CHARS = 120


class UnsupportedFormatError(ValueError):
    pass


def detect_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def load_bytes(data: bytes, filename: str) -> ParsedDocument:
    """Parse `data` into a structured document based on its extension."""
    ext = detect_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(f"Unsupported extension: {ext!r}")
    return _LOADERS[ext](data)


def _is_contents(section: tuple[str, ...]) -> bool:
    return bool(section) and section[0].strip().lower() in _CONTENTS_TITLES


class _SectionBuilder:
    """Accumulates lines under a heading stack and emits segments."""

    def __init__(self) -> None:
        self.segments: list[Segment] = []
        self._stack: list[tuple[int, str]] = []
        self._buffer: list[str] = []

    @property
    def section(self) -> tuple[str, ...]:
        return tuple(title for _, title in self._stack)

    def add_text(self, text: str) -> None:
        if text.strip():
            self._buffer.append(text)

    def flush(self, *, page: int | None = None) -> None:
        if not self._buffer:
            return
        body = clean_text("\n".join(self._buffer))
        self._buffer.clear()
        section = self.section
        if body and not _is_contents(section):
            self.segments.append(Segment(text=body, page=page, section=section))

    def heading(self, level: int, title: str, *, page: int | None = None) -> None:
        self.flush(page=page)
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, title.strip()))

    def extend_heading(self, more: str) -> None:
        """Continue the current heading - PDF headings wrap onto a second line."""
        if not self._stack:
            return
        level, title = self._stack[-1]
        self._stack[-1] = (level, f"{title} {more.strip()}".strip())


# --- plain formats ----------------------------------------------------


def _load_txt(data: bytes) -> ParsedDocument:
    text = clean_text(data.decode("utf-8", errors="replace"))
    return ParsedDocument(segments=[Segment(text=text)] if text else [])


def _load_md(data: bytes) -> ParsedDocument:
    text = clean_text(data.decode("utf-8", errors="replace"))
    if not text:
        return ParsedDocument()

    builder = _SectionBuilder()
    title: str | None = None
    for line in text.split("\n"):
        heading = markdown_heading(line)
        if heading is None:
            builder.add_text(line)
            continue
        level, heading_title = heading
        if title is None and level == 1:
            title = heading_title
        builder.heading(level, heading_title)
    builder.flush()
    return ParsedDocument(segments=builder.segments, title=title)


def _load_json(data: bytes) -> ParsedDocument:
    raw = data.decode("utf-8", errors="replace")
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        text = clean_text(raw)
        return ParsedDocument(segments=[Segment(text=text)] if text else [])

    def dump(value: Any) -> str:
        return clean_text(json.dumps(value, indent=2, ensure_ascii=False))

    if isinstance(parsed, list):
        segments = [
            Segment(text=dump(item), section=(f"Item {i + 1}",))
            for i, item in enumerate(parsed[:10_000])
        ]
        return ParsedDocument(segments=[s for s in segments if s.text])
    if isinstance(parsed, dict):
        segments = [
            Segment(text=dump({key: value}), section=(str(key),))
            for key, value in list(parsed.items())[:10_000]
        ]
        return ParsedDocument(segments=[s for s in segments if s.text])
    return ParsedDocument(segments=[Segment(text=dump(parsed))])


# --- PDF --------------------------------------------------------------


def _pdf_lines(doc: Any) -> list[tuple[int, str, float]]:
    """Every line of the document as (page, text, type size)."""
    lines: list[tuple[int, str, float]] = []
    for page_index, page in enumerate(doc):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s["text"].strip()]
                if not spans:
                    continue
                text = "".join(s["text"] for s in spans).strip()
                if text:
                    lines.append((page_index + 1, text, round(max(s["size"] for s in spans), 1)))
    return lines


def _pdf_heading_levels(lines: list[tuple[int, str, float]]) -> dict[float, int]:
    """Map type size to heading level.

    A PDF carries no headings, only type. The size that sets most of the
    characters is the body; anything meaningfully larger is a heading, and the
    larger it is the higher it sits. This is the document telling us its own
    structure rather than a regular expression guessing at it.
    """
    weight: dict[float, int] = {}
    for _, text, size in lines:
        weight[size] = weight.get(size, 0) + len(text)
    if not weight:
        return {}
    body = max(weight, key=lambda size: weight[size])
    bigger = sorted((size for size in weight if size > body * 1.1), reverse=True)
    return {size: level for level, size in enumerate(bigger[:6], start=1)}


def _load_pdf(data: bytes) -> ParsedDocument:
    import pymupdf  # noqa: PLC0415  (heavy optional dep, loaded on demand)

    # pymupdf ships py.typed but leaves `open`, iteration and `close` unannotated
    doc = pymupdf.open(stream=data, filetype="pdf")  # type: ignore[no-untyped-call]
    try:
        lines = _pdf_lines(doc)
        page_count = doc.page_count
        meta_title = (doc.metadata or {}).get("title") or None
    finally:
        doc.close()  # type: ignore[no-untyped-call]

    levels = _pdf_heading_levels(lines)
    builder = _SectionBuilder()
    current_page = lines[0][0] if lines else None
    previous_heading_size: float | None = None

    for page, text, size in lines:
        if page != current_page:
            builder.flush(page=current_page)
            current_page = page

        level = levels.get(size)
        if level is not None and len(text) <= _MAX_HEADING_CHARS:
            # A heading that wrapped onto a second line is still one heading.
            if previous_heading_size == size:
                builder.extend_heading(text)
            else:
                builder.heading(level, text, page=page)
                previous_heading_size = size
            continue

        previous_heading_size = None
        builder.add_text(text)

    builder.flush(page=current_page)
    return ParsedDocument(
        segments=builder.segments,
        title=meta_title.strip() if isinstance(meta_title, str) and meta_title.strip() else None,
        page_count=page_count,
    )


# --- office formats ---------------------------------------------------


def _docx_heading_level(style_name: str) -> int | None:
    """Word calls them "Heading 1"; anything else is body text."""
    parts = style_name.strip().split()
    if len(parts) == 2 and parts[0].lower() == "heading" and parts[1].isdigit():
        return min(int(parts[1]), 6)
    if style_name.strip().lower() == "title":
        return 1
    return None


def _load_docx(data: bytes) -> ParsedDocument:
    from docx import Document  # noqa: PLC0415  (heavy optional dep, loaded on demand)
    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    doc = Document(io.BytesIO(data))
    builder = _SectionBuilder()
    body = doc.element.body

    # Walking the body's own children keeps tables in reading order. Iterating
    # `doc.paragraphs` instead drops every table on the floor.
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            level = _docx_heading_level(paragraph.style.name if paragraph.style else "")
            if level is not None and paragraph.text.strip():
                builder.heading(level, paragraph.text)
            else:
                builder.add_text(paragraph.text)
        elif tag == "tbl":
            table = Table(child, doc)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    builder.add_text(" | ".join(cells))

    builder.flush()
    return ParsedDocument(segments=builder.segments)


def _load_xlsx(data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook  # noqa: PLC0415  (heavy optional dep)

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    segments: list[Segment] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
        text = clean_text("\n".join(rows))
        if text:
            segments.append(Segment(text=text, section=(sheet.title,)))
    return ParsedDocument(segments=segments)


# --- HTML -------------------------------------------------------------

_HTML_NOISE: Final[tuple[str, ...]] = (
    "script",
    "style",
    "noscript",
    "iframe",
    "svg",
    "form",
    "nav",
    "header",
    "footer",
    "aside",
)
_HTML_HEADINGS: Final[frozenset[str]] = frozenset({"h1", "h2", "h3", "h4", "h5", "h6"})
_HTML_CONTAINERS: Final[frozenset[str]] = frozenset(
    {
        "html",
        "body",
        "main",
        "article",
        "section",
        "div",
        "ul",
        "ol",
        "dl",
        "figure",
        "blockquote",
        "details",
        "template",
    }
)


def _html_walk(node: Any, builder: _SectionBuilder) -> None:
    for child in node.iter(include_text=False):
        tag = child.tag
        if tag in _HTML_HEADINGS:
            title = child.text(deep=True, separator=" ", strip=True)
            if title:
                builder.heading(int(tag[1]), title)
            continue
        if tag == "table":
            for row in child.css("tr"):
                cells = [
                    cell.text(deep=True, separator=" ", strip=True) for cell in row.css("th, td")
                ]
                if any(cells):
                    builder.add_text(" | ".join(cells))
            continue
        if tag in _HTML_CONTAINERS and child.css_first(
            "h1, h2, h3, h4, h5, h6, div, section, article, ul, ol, table, p"
        ):
            # Descend only while there is still structure below. A container
            # holding nothing but inline markup is a paragraph, and taking its
            # text whole keeps the sentence together.
            _html_walk(child, builder)
            continue
        text = child.text(deep=True, separator=" ", strip=True)
        if text:
            builder.add_text(text)


def parse_html(html: str) -> tuple[ParsedDocument, str | None]:
    """Parse a page into sections. Returns the document and its `<title>`."""
    from selectolax.parser import HTMLParser  # noqa: PLC0415  (heavy optional dep)

    tree = HTMLParser(html)
    for selector in _HTML_NOISE:
        for node in tree.css(selector):
            node.decompose()

    title_node = tree.css_first("title")
    title = title_node.text(strip=True) if title_node else None

    # A page that declares its own main region is parsed from there down, so
    # chrome wearing no landmark tag stops opening every chunk.
    root = tree.css_first("main") or tree.css_first("article") or tree.body
    builder = _SectionBuilder()
    if root is not None:
        _html_walk(root, builder)
    builder.flush()
    return ParsedDocument(segments=builder.segments, title=title), title


def _load_html(data: bytes) -> ParsedDocument:
    document, _ = parse_html(data.decode("utf-8", errors="replace"))
    return document


_LOADERS: Final[dict[str, Callable[[bytes], ParsedDocument]]] = {
    "txt": _load_txt,
    "md": _load_md,
    "json": _load_json,
    "pdf": _load_pdf,
    "docx": _load_docx,
    "xlsx": _load_xlsx,
    "html": _load_html,
    "htm": _load_html,
}
